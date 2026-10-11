#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PPTX 产物保真 + 结构校验（与 validate_docx.py 同口径）。

用法：
  python validate_pptx.py --json <中间json> --pptx <产物.pptx>

判定：字符级 CJK 覆盖率 ≥99%（幻灯片正文 + 表格 + 演讲者备注合计，
因汇报版式"要点上页、正文进备注"，正文必须计入备注）。结构检查：
页数、空页数、表格数、图像数、截断告警。
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

CJK = re.compile(r"[\u4e00-\u9fff]")
EMOJI = re.compile(
    "[\U0001F000-\U0001FAFF\u2600-\u27BF\uFE0F\u2190-\u21FF\u2B00-\u2BFF]")


def norm_chars(s):
    s = EMOJI.sub("", s or "")
    return [c for c in s if CJK.match(c)]


def source_markdown(data):
    if data.get("sources") or data.get("analysis"):
        parts = []
        ana = data.get("analysis")
        if ana:
            parts.append(ana.get("markdown") if isinstance(ana, dict) else str(ana))
        parts.extend(s.get("markdown") or "" for s in data.get("sources") or [])
        return "\n\n".join(parts)
    return data.get("markdown") or ""


def deck_text(path):
    """收集全部文本：形状文本 + 表格单元格 + 演讲者备注。"""
    from pptx import Presentation
    prs = Presentation(path)
    texts, n_tables, n_pics, empty = [], 0, 0, 0
    for slide in prs.slides:
        page_has_text = False
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text
                if t.strip():
                    texts.append(t)
                    page_has_text = True
            if getattr(shape, "has_table", False) and shape.has_table:
                n_tables += 1
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            texts.append(cell.text)
                            page_has_text = True
            if shape.shape_type == 13:  # PICTURE
                n_pics += 1
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text
            if nt.strip():
                texts.append(nt)
                page_has_text = True
        if not page_has_text:
            empty += 1
    return len(prs.slides), texts, n_tables, n_pics, empty


def main():
    ap = argparse.ArgumentParser(description="omni2docx PPTX 保真校验器")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--json", help="渲染时输入的中间 JSON")
    g.add_argument("--md", help="--md 直出时的源 markdown（替代 --json 的源）")
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--max-slides", type=int, default=None,
                    help="渲染时传过的 --max-slides（用于判断截断）")
    args = ap.parse_args()

    if args.md:
        src_md = open(args.md, encoding="utf-8").read()
    else:
        src_md = source_markdown(json.load(open(args.json, encoding="utf-8")))

    n_slides, texts, n_tables, n_pics, empty = deck_text(args.pptx)
    out_text = "\n".join(texts)

    src_cjk = Counter(norm_chars(src_md))
    out_cjk = Counter(norm_chars(out_text))
    if src_cjk:
        hit = sum(min(v, out_cjk.get(k, 0)) for k, v in src_cjk.items())
        coverage = hit / sum(src_cjk.values())
        missing = {k: v for k, v in src_cjk.items() if out_cjk.get(k, 0) < v}
    else:
        coverage, missing = 1.0, {}

    print("=" * 64)
    print("Omni2Docx PPTX 保真 + 结构校验 :", os.path.basename(args.pptx))
    print("=" * 64)
    print("[内容保真]（正文+表格+演讲者备注合并口径）")
    print("  源文 CJK 字符种类 :", len(src_cjk))
    print("  deck 含源文 CJK 覆盖率 : %.2f%%" % (coverage * 100))
    print("  缺失 CJK 字符(种类) :", len(missing), list(missing)[:20])
    print("[结构]")
    print("  总页数 :", n_slides, ("（达 --max-slides %d 上限，已截断）" % args.max_slides)
          if args.max_slides and n_slides >= args.max_slides else "")
    print("  空页数 :", empty)
    print("  表格数 :", n_tables)
    print("  图像数 :", n_pics)

    problems = []
    if coverage < 0.99:
        problems.append("CJK 覆盖率 <99%")
    if n_slides == 0:
        problems.append("零页 deck")
    if empty > max(1, n_slides // 10):
        problems.append("空页过多（%d/%d）" % (empty, n_slides))
    print("判定 :", "PASS" if not problems else "FAIL（%s）" % "；".join(problems))
    sys.exit(0 if not problems else 1)


if __name__ == "__main__":
    main()
