#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""run_matrix_test.py — omni2docx docx 引擎真实语料矩阵回归 + XML 级专项断言。

读取 tests/fixtures/ 里缓存的真实 omni 解析产物（多源类型 markdown）+ 合成边界
fixture，批量跑 scripts/build_docx.py（多 profile / 开关组合），
scripts/validate_docx.py 保真校验，再做 XML 级断言：PAGE 域、emoji 剥离、
引用块楷体、递归 inline。

路径全部相对本文件，跨平台可跑（python3 即可，需装 python-docx）。
生成产物落在 tests/out/（不入库，见 .gitignore）。
"""
import json
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
FIX = HERE / "fixtures"
OUT = HERE / "out"
SKILL = HERE.parent / "scripts"
PY = sys.executable
BUILD = str(SKILL / "build_docx.py")
VALID = str(SKILL / "validate_docx.py")

EMOJI_RE = re.compile(
    "[\U0001F000-\U0001FAFF\U00002600-\U000027BF\U0001F1E6-\U0001F1FF"
    "\u2B00-\u2BFF\uFE0F\u2705\u274C\u2757\u3030\u303D\u3297\u3299]"
)

# (case_id, fixture md 文件名, profile, extra args, 断言集)
CASES = [
    ("M01_etf_xlsx_gov", "m01_etf_xlsx.omni.md", "gov", ["--page-numbers"], ["page_field"]),
    ("M02_bond_pdf_finance", "m02_bond_pdf.omni.md", "finance", [], []),
    ("M03_report_pdf_legal", "m03_report_pdf.omni.md", "legal", ["--strip-emoji", "--page-numbers"], ["page_field", "no_emoji"]),
    ("M04_html_academic", "m04_html.omni.md", "academic", [], []),
    ("M05_audio_default", "m05_audio.omni.md", "default", [], []),
    ("M06_docx_custom_profile", "m06_docx.omni.md",
     '{"body_east":"仿宋","body_latin":"Times New Roman","body_size":14,"line_spacing":"28pt","first_line_indent":2,"table_style":"Table Grid","page_number":true,"heading_east":"黑体"}',
     [], ["page_field"]),
    ("M07_pptx_default", "m07_pptx.omni.md", "default", [], []),
    ("M08_jpg_ocr_strip", "m08_jpg_ocr.omni.md", "default", ["--strip-emoji"], ["no_emoji"]),
    ("M09_huge1000p_gov", "m09_huge1000p.omni.md", "gov", ["--page-numbers"], ["page_field"]),
    ("M10_truncated_default", "m10_truncated.omni.md", "default", [], []),
    ("M11_edge_all_flags", "edge_case.md", "academic", ["--page-numbers", "--strip-emoji"], ["page_field", "no_emoji", "quote_kaiti", "nested_inline"]),
    ("M12_edge_default_noStrip", "edge_case.md", "default", [], ["emoji_kept", "nested_inline"]),
    ("M13_ocr_spacing", "m11_ocr_spacing.omni.md", "default", [], ["cjk_space_clean", "latin_space_preserved"]),
]

results = []


def wrap_json(md_path: Path, tmp_json: Path):
    md = md_path.read_text(encoding="utf-8")
    data = {"markdown": md, "grounding": None, "outline": None,
            "source": md_path.name, "detail": "text"}
    tmp_json.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return tmp_json


def docx_text(docx_path: Path) -> str:
    from docx import Document
    doc = Document(str(docx_path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def find_quote_para(docx_path: Path):
    from docx import Document
    doc = Document(str(docx_path))
    for p in doc.paragraphs:
        if "这是一段引用" in p.text:
            return p
    return None


def assert_flags(flag: str, docx_path: Path) -> str:
    import zipfile
    text = docx_text(docx_path)
    from docx import Document
    doc = Document(str(docx_path))
    if flag == "page_field":
        z = zipfile.ZipFile(str(docx_path))
        for n in z.namelist():
            if "footer" in n and n.endswith(".xml"):
                x = z.read(n).decode("utf-8", "ignore")
                if "fldSimple" in x and "PAGE" in x:
                    return "PASS"
        return "FAIL(footer part 无 PAGE 域)"
    if flag == "no_emoji":
        m = EMOJI_RE.findall(text)
        return "PASS" if not m else f"FAIL(残留 emoji {m[:5]})"
    if flag == "emoji_kept":
        m = EMOJI_RE.findall(text)
        return "PASS" if m else "FAIL(emoji 应保留但为零)"
    if flag == "quote_kaiti":
        p = find_quote_para(docx_path)
        if p is None:
            return "FAIL(找不到引用段)"
        fonts = set()
        for r in p.runs:
            rPr = r._element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
            if rPr is not None:
                rf = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                if rf is not None:
                    fonts.add(rf.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia'))
        return "PASS" if "楷体" in fonts else f"FAIL(引用段 eastAsia={fonts})"
    if flag == "nested_inline":
        def norm_t(s):
            return re.sub(r"[\s*]+", "", s)
        for p in doc.paragraphs:
            if "加粗中包含" in p.text:
                bold_runs = [r.text for r in p.runs if r.bold]
                joined = norm_t("".join(bold_runs))
                need = ["加粗中包含", "code片段", "链接", "斜体"]
                missing = [w for w in need if w not in joined]
                if not missing and "*" not in joined:
                    return "PASS"
                return f"FAIL(粗体 run 缺 {missing} 或残留 *)"
        return "FAIL(找不到嵌套行内段)"
    if flag == "cjk_space_clean":
        if "他 们" in text:
            return "FAIL(CJK 逐字空格未被清理)"
        if "他们讨论" not in text:
            return "FAIL(清理后 CJK 正文缺失)"
        return "PASS"
    if flag == "latin_space_preserved":
        if "W e d i s c u s s e d" in text:
            return "PASS"
        return "FAIL(拉丁 OCR 拆字被意外合并/丢失)"
    return "SKIP"


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for case_id, md_name, profile, extra, asserts in CASES:
        md_path = FIX / md_name
        rec = {"case": case_id, "md": md_path.name,
               "profile": (profile[:60] + "…") if len(profile) > 60 else profile,
               "args": extra}
        if not md_path.exists() or md_path.stat().st_size == 0:
            rec.update(status="SKIP_EMPTY_MD", validate="-", asserts={})
            results.append(rec)
            print(f"[SKIP] {case_id}: 源 md 缺失或为空")
            continue
        tmp_json = OUT / f"_t_{case_id}.json"
        wrap_json(md_path, tmp_json)
        docx_out = OUT / f"t_{case_id}.docx"
        cmd = [PY, BUILD, "--json", str(tmp_json), "--out", str(docx_out),
               "--profile", profile] + extra
        r = subprocess.run(cmd, capture_output=True, text=True,
                           encoding="utf-8", errors="replace")
        rec["build_rc"] = r.returncode
        if r.returncode != 0:
            rec.update(status="BUILD_FAIL", validate="-", asserts={},
                       err=(r.stderr or r.stdout)[-400:])
            results.append(rec)
            print(f"[BUILD FAIL] {case_id}: {rec['err']}")
            continue
        v = subprocess.run([PY, VALID, "--json", str(tmp_json), "--docx", str(docx_out)],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace")
        val_ok = v.returncode == 0
        cov = re.search(r"覆盖率 : ([\d.]+)%", v.stdout)
        rec["validate"] = ("PASS" if val_ok else "CHECK") + \
            (f" coverage={cov.group(1)}%" if cov else "")
        rec["asserts"] = {f: assert_flags(f, docx_out) for f in asserts}
        rec["status"] = "OK" if val_ok and all(a == "PASS"
                          for a in rec["asserts"].values()) else "ISSUE"
        results.append(rec)
        print(f"[{rec['status']}] {case_id} | {rec['validate']} | {rec['asserts']}")

    summary_path = OUT / "matrix_test_report.json"
    summary_path.write_text(json.dumps(results, ensure_ascii=False, indent=2),
                           encoding="utf-8")
    print("\n=== 汇总 ===")
    bad = 0
    for rec in results:
        ok = rec["status"] == "OK"
        bad += 0 if ok else 1
        print(rec["case"], "->", rec["status"], "|", rec.get("validate", "-"),
              "|", rec.get("asserts", {}))
    print(f"合计 {len(results)} 例，OK {len(results) - bad}，ISSUE/SKIP {bad}")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
